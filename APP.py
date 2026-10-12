import io
import random
import re
from collections import Counter
from datetime import datetime

import dash
from dash import Dash, Input, Output, State, dcc, html, dash_table, no_update
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.express as px
import requests
from bs4 import BeautifulSoup

# ============================================================
# 1. 上網抓取資料
# ============================================================
def fetch_draws(max_draws=60):
    url = "https://lottery.hk/en/mark-six/results/2026"
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                      "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }
    try:
        resp = requests.get(url, headers=headers, timeout=15)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")

        draws = []
        for tr in soup.select("table tr"):
            tds = tr.find_all("td")
            if len(tds) < 3:
                continue
            draw_no = tds[0].get_text(strip=True)
            date_str = tds[1].get_text(strip=True)
            balls_text = tds[2].get_text(separator=" ", strip=True)
            nums = [int(x) for x in re.findall(r"\d+", balls_text)]

            if len(nums) >= 7 and "/" in draw_no:
                try:
                    dt = datetime.strptime(date_str, "%d/%m/%Y")
                    formatted_date = dt.strftime("%Y/%m/%d")
                except Exception:
                    formatted_date = date_str

                main_nums = sorted(nums[:6])
                special = nums[6]
                draws.append((formatted_date, main_nums, special))

                if len(draws) >= max_draws:
                    break

        if len(draws) < 20:
            raise ValueError(f"只抓到 {len(draws)} 期，資料不足")

        print(f"✅ 成功抓取最新 {len(draws)} 期資料")
        return draws

    except Exception as e:
        print(f"⚠️ 抓取線上資料失敗：{e}")
        return None


FALLBACK_DRAWS = [
    ("2026/10/10", [12, 14, 20, 33, 34, 37], 31),
    ("2026/10/08", [18, 19, 24, 31, 40, 44], 5),
    ("2026/10/06", [9, 29, 31, 34, 36, 41], 28),
    ("2026/10/03", [8, 23, 29, 31, 34, 43], 48),
    ("2026/09/26", [4, 28, 31, 44, 47, 48], 19),
    ("2026/09/22", [2, 34, 35, 43, 45, 46], 41),
    ("2026/09/19", [5, 8, 11, 14, 19, 31], 21),
    ("2026/09/17", [4, 17, 31, 41, 47, 48], 8),
    ("2026/09/15", [2, 5, 6, 9, 29, 30], 19),
    ("2026/09/12", [9, 18, 24, 33, 40, 47], 11),
    ("2026/09/10", [22, 23, 31, 32, 35, 45], 34),
    ("2026/09/08", [9, 23, 28, 29, 35, 41], 38),
    ("2026/09/05", [9, 18, 26, 30, 33, 45], 28),
    ("2026/08/29", [4, 7, 8, 11, 26, 30], 42),
    ("2026/08/27", [12, 17, 25, 31, 44, 45], 35),
    ("2026/08/25", [1, 18, 19, 25, 34, 38], 7),
    ("2026/08/22", [7, 9, 12, 25, 34, 40], 29),
    ("2026/08/20", [7, 9, 10, 15, 24, 46], 34),
    ("2026/08/18", [7, 8, 9, 14, 39, 41], 49),
    ("2026/08/15", [4, 16, 25, 27, 28, 33], 14),
    ("2026/08/13", [21, 27, 35, 40, 47, 48], 23),
    ("2026/08/11", [11, 19, 28, 36, 43, 48], 13),
    ("2026/08/08", [1, 3, 7, 31, 32, 49], 47),
    ("2026/08/06", [13, 24, 29, 30, 42, 48], 3),
    ("2026/08/04", [4, 11, 13, 16, 31, 33], 38),
    ("2026/08/01", [1, 7, 16, 22, 32, 37], 23),
    ("2026/07/30", [1, 2, 14, 17, 23, 35], 48),
    ("2026/07/28", [4, 7, 14, 20, 21, 30], 34),
    ("2026/07/25", [5, 7, 12, 26, 31, 42], 11),
    ("2026/07/23", [4, 7, 15, 29, 32, 46], 23),
    ("2026/07/21", [1, 6, 37, 41, 45, 47], 8),
    ("2026/07/16", [18, 21, 26, 30, 32, 44], 27),
    ("2026/07/14", [10, 16, 19, 27, 31, 44], 17),
    ("2026/07/11", [2, 5, 7, 11, 41, 46], 43),
    ("2026/07/09", [1, 7, 18, 21, 33, 38], 28),
    ("2026/07/07", [5, 34, 37, 43, 48, 49], 27),
    ("2026/07/05", [3, 14, 19, 20, 31, 38], 44),
    ("2026/07/02", [5, 17, 26, 31, 34, 46], 43),
    ("2026/06/30", [12, 14, 17, 18, 22, 40], 13),
    ("2026/06/28", [5, 7, 26, 32, 37, 40], 35),
    ("2026/06/25", [12, 16, 19, 20, 28, 30], 29),
    ("2026/06/23", [2, 15, 25, 34, 39, 41], 10),
    ("2026/06/20", [5, 7, 10, 15, 16, 41], 46),
    ("2026/06/18", [6, 20, 26, 31, 34, 47], 24),
    ("2026/06/14", [10, 12, 17, 29, 32, 35], 5),
    ("2026/06/11", [3, 8, 15, 22, 36, 49], 41),
    ("2026/06/09", [1, 9, 18, 27, 33, 42], 16),
    ("2026/06/06", [4, 13, 21, 28, 39, 45], 7),
    ("2026/06/04", [2, 11, 19, 24, 38, 46], 30),
    ("2026/06/02", [6, 14, 23, 31, 40, 48], 9),
]


def process_draws(all_draws):
    draws_for_table = all_draws[:20]
    draws_for_stats = all_draws[:50]

    table_rows = []
    for date_str, nums, extra in draws_for_table:
        table_rows.append({
            "date": date_str,
            "n1": nums[0], "n2": nums[1], "n3": nums[2],
            "n4": nums[3], "n5": nums[4], "n6": nums[5],
            "s_num": extra,
        })
    df_draws = pd.DataFrame(table_rows)

    all_main_numbers = []
    for _, nums, _ in draws_for_stats:
        all_main_numbers.extend(nums)

    counts = Counter(all_main_numbers)
    df_counts = pd.DataFrame(
        [{"號碼": f"{num:02d}", "出現次數": counts.get(num, 0)} for num in range(1, 50)]
    )
    df_top10 = df_counts.sort_values(by="出現次數", ascending=False).head(10)

    cold_nums = [f"{num:02d}" for num in range(1, 50) if counts.get(num, 0) == 0]
    cold_sorted = df_counts.sort_values(by="出現次數").head(10)
    cold_list_text = ", ".join(
        [f"{row['號碼']}({row['出現次數']})" for _, row in cold_sorted.iterrows()]
    )

    fig = px.bar(
        df_top10, x="號碼", y="出現次數",
        title="🔥 近 50 期熱門號碼 Top 10（主號）",
        text="出現次數", color="出現次數", color_continuous_scale="Reds",
    )
    fig.update_layout(xaxis_type="category", template="plotly_white", margin=dict(l=20, r=20, t=50, b=20))

    cold_zero_text = ", ".join(cold_nums) if cold_nums else "無（所有號碼都有出現）"
    update_time = datetime.now().strftime("%Y-%m-%d %H:%M")

    return {
        "table_data": df_draws.to_dict("records"),
        "figure": fig,
        "cold_zero_text": cold_zero_text,
        "cold_list_text": cold_list_text,
        "update_time": update_time,
        "n_table": len(draws_for_table),
        "n_stats": len(draws_for_stats),
        "history_draws_20": draws_for_table,
        "all_main_numbers_50": all_main_numbers,
    }


_initial_raw = fetch_draws(60) or FALLBACK_DRAWS
_initial_data = process_draws(_initial_raw)

# ============================================================
# 2. Dash App
# ============================================================
app = Dash(__name__)
app.title = "六合彩 Dashboard"
server = app.server

app.layout = html.Div(
    style={"backgroundColor": "#f4f6f9", "fontFamily": "Microsoft JhengHei, sans-serif", "padding": "20px"},
    children=[
        dcc.Store(id="store-data", data=_initial_data),

        html.Div(
            style={
                "backgroundColor": "#1F497D", "color": "white", "padding": "15px 20px",
                "borderRadius": "8px", "marginBottom": "20px", "display": "flex",
                "justifyContent": "space-between", "alignItems": "center", "flexWrap": "wrap", "gap": "10px",
            },
            children=[
                html.Div([
                    html.H2("🎯 香港六合彩 (Mark Six) 數據分析 Dashboard", style={"margin": 0}),
                    html.P(
                        id="subtitle-text",
                        children=f"表格：最近 {_initial_data['n_table']} 期 ｜ 熱門/冷門統計：最近 {_initial_data['n_stats']} 期 ｜ 更新時間：{_initial_data['update_time']}",
                        style={"margin": "5px 0 0 0", "fontSize": "14px", "opacity": 0.9},
                    ),
                ]),
                html.Button(
                    "🔄 重新整理資料", id="btn-refresh", n_clicks=0,
                    style={"backgroundColor": "#ffc107", "color": "#212529", "border": "none",
                           "padding": "10px 18px", "fontSize": "15px", "borderRadius": "5px",
                           "cursor": "pointer", "fontWeight": "bold"},
                ),
            ],
        ),

        html.Div(
            style={"backgroundColor": "white", "padding": "20px", "borderRadius": "8px",
                   "boxShadow": "0 2px 4px rgba(0,0,0,0.1)", "marginBottom": "20px"},
            children=[
                html.H4("🎲 1 - 49 隨機攪珠選號", style={"marginTop": 0}),
                html.Button(
                    "隨機產生 6 個號碼", id="btn-generate", n_clicks=0,
                    style={"backgroundColor": "#28a745", "color": "white", "border": "none",
                           "padding": "10px 20px", "fontSize": "16px", "borderRadius": "5px", "cursor": "pointer"},
                ),
                html.Div(id="gen-result", style={"marginTop": "15px", "fontSize": "20px", "fontWeight": "bold", "color": "#28a745"}),
            ],
        ),

        html.Div(
            style={"display": "flex", "gap": "20px", "marginBottom": "20px", "flexWrap": "wrap"},
            children=[
                html.Div(
                    style={"flex": "2", "backgroundColor": "white", "padding": "15px", "borderRadius": "8px",
                           "boxShadow": "0 2px 4px rgba(0,0,0,0.1)", "minWidth": "350px"},
                    children=[dcc.Graph(id="hot-chart", figure=_initial_data["figure"])],
                ),
                html.Div(
                    style={"flex": "1", "backgroundColor": "white", "padding": "20px", "borderRadius": "8px",
                           "boxShadow": "0 2px 4px rgba(0,0,0,0.1)", "minWidth": "280px"},
                    children=[
                        html.H4("❄️ 冷門號碼分析（近 50 期）", style={"color": "#0275d8", "marginTop": 0}),
                        html.P("完全未開出（0 次）：", style={"marginBottom": "4px"}),
                        html.H4(id="cold-zero", children=_initial_data["cold_zero_text"], style={"color": "#d9534f", "marginTop": 0}),
                        html.Hr(),
                        html.P("出現次數最少的 10 個號碼：", style={"marginBottom": "4px"}),
                        html.P(id="cold-list", children=_initial_data["cold_list_text"],
                               style={"color": "#5bc0de", "fontSize": "15px", "fontWeight": "bold"}),
                    ],
                ),
            ],
        ),

        html.Div(
            style={"backgroundColor": "white", "padding": "20px", "borderRadius": "8px",
                   "boxShadow": "0 2px 4px rgba(0,0,0,0.1)", "marginBottom": "20px"},
            children=[
                html.Div(
                    style={"display": "flex", "justifyContent": "space-between", "alignItems": "center",
                           "marginBottom": "15px", "flexWrap": "wrap", "gap": "10px"},
                    children=[
                        html.H4("📋 近 20 期詳細開獎紀錄", style={"margin": 0}),
                        html.Button(
                            "📊 下載 Excel 分析報表", id="btn-excel", n_clicks=0,
                            style={"backgroundColor": "#17a2b8", "color": "white", "border": "none",
                                   "padding": "10px 18px", "fontSize": "14px", "borderRadius": "5px", "cursor": "pointer"},
                        ),
                        dcc.Download(id="download-excel"),
                    ],
                ),
                dash_table.DataTable(
                    id="history-table",
                    data=_initial_data["table_data"],
                    columns=[
                        {"name": "期數/日期", "id": "date"},
                        {"name": "號碼 1", "id": "n1"}, {"name": "號碼 2", "id": "n2"},
                        {"name": "號碼 3", "id": "n3"}, {"name": "號碼 4", "id": "n4"},
                        {"name": "號碼 5", "id": "n5"}, {"name": "號碼 6", "id": "n6"},
                        {"name": "特別號", "id": "s_num"},
                    ],
                    style_header={"backgroundColor": "#1F497D", "color": "white", "fontWeight": "bold", "textAlign": "center"},
                    style_cell={"textAlign": "center", "padding": "8px"},
                    style_data_conditional=[
                        {"if": {"row_index": "odd"}, "backgroundColor": "#f9fbfd"},
                        {"if": {"column_id": "s_num"}, "color": "#C00000", "fontWeight": "bold"},
                    ],
                ),
            ],
        ),
    ],
)

# ============================================================
# 3. Callbacks
# ============================================================
@app.callback(
    Output("store-data", "data"),
    Output("subtitle-text", "children"),
    Output("hot-chart", "figure"),
    Output("cold-zero", "children"),
    Output("cold-list", "children"),
    Output("history-table", "data"),
    Input("btn-refresh", "n_clicks"),
    prevent_initial_call=True,
)
def refresh_data(n_clicks):
    if not n_clicks:
        return no_update, no_update, no_update, no_update, no_update, no_update

    new_raw = fetch_draws(60) or FALLBACK_DRAWS
    data = process_draws(new_raw)

    subtitle = (
        f"表格：最近 {data['n_table']} 期 ｜ "
        f"熱門/冷門統計：最近 {data['n_stats']} 期 ｜ "
        f"更新時間：{data['update_time']}"
    )
    return data, subtitle, data["figure"], data["cold_zero_text"], data["cold_list_text"], data["table_data"]


@app.callback(Output("gen-result", "children"), Input("btn-generate", "n_clicks"))
def generate_random_numbers(n_clicks):
    if n_clicks == 0:
        return "點擊按鈕開始選號..."
    nums = sorted(random.sample(range(1, 50), 6))
    return f"抽中號碼：[ {'  ·  '.join([f'{n:02d}' for n in nums])} ]"


@app.callback(
    Output("download-excel", "data"),
    Input("btn-excel", "n_clicks"),
    State("store-data", "data"),
    prevent_initial_call=True,
)
def export_excel(n_clicks, store_data):
    if not store_data:
        return no_update

    history_draws = store_data["history_draws_20"]
    all_main_numbers = store_data["all_main_numbers_50"]

    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "統計摘要"
    ws_detail = wb.create_sheet(title="近20期詳細紀錄")

    font_title = Font(name="Microsoft JhengHei", size=15, bold=True, color="1F497D")
    font_header = Font(name="Microsoft JhengHei", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Microsoft JhengHei", size=11)
    font_bold = Font(name="Microsoft JhengHei", size=11, bold=True)
    font_special = Font(name="Microsoft JhengHei", size=11, bold=True, color="C00000")

    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_top3 = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"),
    )
    align_center = Alignment(horizontal="center", vertical="center")

    ws_detail.cell(row=1, column=1, value="香港六合彩 - 近20期詳細開獎紀錄").font = font_title
    headers_detail = ["期數/日期", "號碼 1", "號碼 2", "號碼 3", "號碼 4", "號碼 5", "號碼 6", "特別號"]
    for c_idx, text in enumerate(headers_detail, 1):
        cell = ws_detail.cell(row=3, column=c_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    for r_idx, (date_str, nums, extra) in enumerate(history_draws, 4):
        for c_idx, val in enumerate([date_str] + nums + [extra], 1):
            cell = ws_detail.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_special if c_idx == 8 else font_body
            cell.alignment = align_center
            cell.border = thin_border
            if r_idx % 2 == 0:
                cell.fill = fill_zebra

    ws_summary.cell(row=1, column=1, value="香港六合彩 - 近50期出現頻率統計分析").font = font_title
    ws_summary.cell(row=3, column=1, value="【1-49號碼次數統計表（近50期）】").font = font_bold

    for c_idx, text in enumerate(["號碼", "出現次數 (主號)", "出現率"], 1):
        cell = ws_summary.cell(row=4, column=c_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    counts = Counter(all_main_numbers)
    for num in range(1, 50):
        r_idx = num + 4
        cnt = counts.get(num, 0)
        for c, val in zip(
            [ws_summary.cell(row=r_idx, column=1), ws_summary.cell(row=r_idx, column=2), ws_summary.cell(row=r_idx, column=3)],
            [num, cnt, cnt / 50]
        ):
            c.value = val
            c.font = font_body
            c.alignment = align_center
            c.border = thin_border
            if num % 2 == 0:
                c.fill = fill_zebra
        ws_summary.cell(row=r_idx, column=3).number_format = "0.0%"

    ws_summary.cell(row=3, column=5, value="🔥【最熱門號碼 Top 10（近50期）】").font = font_bold
    for c_idx, text in enumerate(["名次", "熱門號碼", "出現次數"], 5):
        cell = ws_summary.cell(row=4, column=c_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    for idx, (num, cnt) in enumerate(counts.most_common(10), 1):
        r_idx = idx + 4
        for c, val in zip(
            [ws_summary.cell(row=r_idx, column=5), ws_summary.cell(row=r_idx, column=6), ws_summary.cell(row=r_idx, column=7)],
            [f"第 {idx} 名", num, cnt]
        ):
            c.value = val
            c.font = font_bold if idx <= 3 else font_body
            c.alignment = align_center
            c.border = thin_border
            if idx <= 3:
                c.fill = fill_top3
            elif idx % 2 == 0:
                c.fill = fill_zebra

    for ws in [ws_summary, ws_detail]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 5, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return dcc.send_bytes(output.getvalue(), filename="MarkSix_Dashboard_Report.xlsx")


if __name__ == "__main__":
    app.run(debug=True)
